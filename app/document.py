import os
from pypdf import PdfReader
from docx import Document
import csv
from openpyxl import load_workbook
from pptx import Presentation


def read_pdf(pdf_path):
    reader = PdfReader(pdf_path)

    pages = []

    for page_number, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text()

        if page_text:
            pages.append({
                "page": page_number,
                "text": page_text
            })

    return pages


def read_txt(txt_path):
    with open(txt_path, "r", encoding="utf-8") as file:
        text = file.read()

    return [
        {
            "page": 1,
            "text": text
        }
    ]

def read_docx(docx_path):
    document = Document(docx_path)

    text = "\n\n".join(
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    )

    return [
        {
            "page": 1,
            "text": text
        }
    ]

def read_csv(csv_path):
    rows = []

    with open(csv_path, "r", encoding="utf-8", newline="") as file:
        reader = csv.reader(file)

        for row in reader:
            if row:
                rows.append(" | ".join(row))

    text = "\n".join(rows)

    return [
        {
            "page": 1,
            "text": text
        }
    ]

def read_xlsx(xlsx_path):
    workbook = load_workbook(
        xlsx_path,
        read_only=True,
        data_only=True
    )

    rows = []

    for worksheet in workbook.worksheets:
        rows.append(f"Sheet: {worksheet.title}")

        for row in worksheet.iter_rows(values_only=True):
            values = [
                str(value)
                for value in row
                if value is not None
            ]

            if values:
                rows.append(" | ".join(values))

    text = "\n".join(rows)

    return [
        {
            "page": 1,
            "text": text
        }
    ]

def read_pptx(pptx_path):
    presentation = Presentation(pptx_path)

    slides = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):
        slide_text = []

        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_text.append(shape.text.strip())

        text = "\n".join(slide_text)

        if text:
            slides.append({
                "page": slide_number,
                "text": text
            })

    return slides


def load_documents(folder_path):
    all_pages = []

    for filename in os.listdir(folder_path):

        file_path = os.path.join(folder_path, filename)

        if filename.lower().endswith(".pdf"):

            pages = read_pdf(file_path)

        elif filename.lower().endswith(".txt"):

            pages = read_txt(file_path)


        elif filename.lower().endswith(".docx"):

            pages = read_docx(file_path)


        elif filename.lower().endswith(".csv"):

            pages = read_csv(file_path)


        elif filename.lower().endswith(".xlsx"):

            pages = read_xlsx(file_path)

        elif filename.lower().endswith(".pptx"):

            pages = read_pptx(file_path)

        else:

            continue


        for page in pages:
            all_pages.append({
                "filename": filename,
                "page": page["page"],
                "text": page["text"]
            })

    return all_pages


def split_text(text, chunk_size=1000):

    text = "\n".join(
        line.strip()
        for line in text.splitlines()
        if line.strip()
    )

    paragraphs = text.split("\n\n")

    chunks = []
    current_chunk = ""

    for paragraph in paragraphs:

        if len(current_chunk) + len(paragraph) <= chunk_size:
            current_chunk += paragraph + "\n\n"

        else:

            if current_chunk.strip():
                chunks.append(current_chunk.strip())

            current_chunk = paragraph + "\n\n"

    if current_chunk.strip():
        chunks.append(current_chunk.strip())

    chunks = [
        chunk
        for chunk in chunks
        if len(chunk) >= 100
    ]

    return chunks


def create_document_chunks(documents):

    all_chunks = []

    for document in documents:

        chunks = split_text(document["text"])

        for chunk in chunks:

            all_chunks.append({
                "text": chunk,
                "source": document["filename"],
                "page": document["page"]
            })

    return all_chunks










